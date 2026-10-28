from pathlib import Path
from docx import Document

root = Path(r"C:\Users\Lenovo\Desktop\省一改")
src = next(root.glob("项目报告书*ProMax.docx"))
doc = Document(src)

print("TOC/FRONT")
for i, p in enumerate(doc.paragraphs[:92]):
    if i >= 15:
        print(i, repr(p.style.style_id), repr(p.text))

print("\nTABLES")
for ti, table in enumerate(doc.tables, 1):
    print("TABLE", ti)
    for row in table.rows:
        print(" | ".join(c.text.replace("\n", " / ") for c in row.cells))

print("\nCH3 MIDDLE")
for i, p in enumerate(doc.paragraphs):
    if 178 <= i <= 225:
        print(i, repr(p.style.style_id), repr(p.text))
