from pathlib import Path
from docx import Document

root = Path(r"C:\Users\Lenovo\Desktop\省一改")
path = root / "项目报告书-智析万卷：面向ESG报告的自适应布局分析与要素提取-专业团队ProMax-国赛专项修订版.docx"
doc = Document(path)

start = next(i for i,p in enumerate(doc.paragraphs) if p.text == "ESG报告数据处理与指标体系构建")
end = next(i for i,p in enumerate(doc.paragraphs) if p.text == "ESGMiner平台设计与实现")
ch = doc.paragraphs[start:end]
text = "\n".join(p.text for p in ch)

print("FILE", path.name, path.stat().st_size)
print("COUNTS paragraphs/tables/shapes", len(doc.paragraphs), len(doc.tables), len(doc.inline_shapes))
print("CHAPTER_RANGE", start, end, "paragraphs", end-start)
print("FORMULA_CAPTIONS", [p.text.strip() for p in ch if "公式" in p.text])
print("FIGURES", [p.text for p in doc.paragraphs if p.text.startswith("图 ") and "\t" not in p.text])
print("TABLE_CAPTIONS", [p.text for p in doc.paragraphs if p.text.startswith("表 ") and "\t" not in p.text])

print("\nHEADINGS")
for p in ch:
    if p.style.style_id in {"a", "a0", "gao"}:
        print(p.style.style_id, p.text)

print("\nREQUIRED")
checks = {
    "flow_caption": "企业端到端交付闭环图" in text,
    "enterprise_table": "企业运行能力与降级策略" in text,
    "batch6": "默认每批6个字段" in text,
    "max12": "最多12次" in text,
    "topk": "Top-K" in text,
    "year2024": "默认以2024年" in text,
    "thresholds": "置信度低于0.7" in text and "证据链得分低于50" in text,
    "tfidf": "字符n-gram TF-IDF余弦检索" in text,
    "json_csv": "API以JSON" in text and "导出中心" in text and "CSV" in text,
    "metrics": all(x in text for x in ["56项", "12项结果正确", "字段串扰12项", "数值错误3项", "年份错误1项", "80.95%"]),
    "limitation": "尚未对系统识别未披露状态的准确性进行独立评价" in text,
    "transition2": "本章完成了数据来源、指标标准、产品处理流程及评估口径的总体设计" in text,
    "transition3": "上述算法与工程机制完成了从PDF输入到结构化结果" in text,
}
for k,v in checks.items():
    print(k, v)

print("\nBANNED")
for bad in ["请求头伪装", "规避动态加载", "用户只需登录", "时间感知阈值", "Blake2b", "384维", "模型性能指标", "无遗漏", "性能完整可靠", "开箱即用"]:
    print(bad, text.count(bad))

print("\nTABLES")
for i,t in enumerate(doc.tables,1):
    print(i, len(t.rows), len(t.columns), " | ".join(c.text for c in t.rows[0].cells))

if not all(checks.values()):
    raise SystemExit("Required content check failed")
