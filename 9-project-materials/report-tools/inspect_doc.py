from pathlib import Path
from docx import Document
from docx.oxml.ns import qn

root = Path(r"C:\Users\Lenovo\Desktop\省一改")
src = next(root.glob("项目报告书*ProMax.docx"))
doc = Document(src)

body = doc.element.body
inside = False
for pos, el in enumerate(list(body)):
    tag = el.tag.rsplit("}", 1)[-1]
    if tag == "p":
        texts = el.xpath(".//w:t/text()")
        text = "".join(texts).strip()
        pstyle = el.find(qn("w:pPr"))
        style = ""
        if pstyle is not None:
            ps = pstyle.find(qn("w:pStyle"))
            if ps is not None:
                style = ps.get(qn("w:val"), "")
        drawings = len(el.xpath(".//w:drawing"))
        maths = len(el.xpath(".//m:oMath | .//m:oMathPara"))
        if text == "ESG报告数据处理与指标体系构建" and style:
            inside = True
        if inside:
            print(f"{pos:03d} P style={style!r} draw={drawings} math={maths} text={text[:130]!r}")
        if inside and text == "产品应用情况":
            break
    elif tag == "tbl" and inside:
        rows = el.findall(qn("w:tr"))
        first = ""
        if rows:
            first = "".join(rows[0].xpath(".//w:t/text()"))
        print(f"{pos:03d} T rows={len(rows)} first={first[:100]!r}")

print("paragraphs", len(doc.paragraphs), "tables", len(doc.tables), "shapes", len(doc.inline_shapes))
print("fields", len(doc.element.xpath(".//w:fldSimple | .//w:instrText")))
