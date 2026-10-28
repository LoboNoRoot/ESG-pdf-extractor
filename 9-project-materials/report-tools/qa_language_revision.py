from pathlib import Path
from docx import Document
import zipfile

root = Path(r"C:\Users\Lenovo\Desktop\省一改")
src = root / "项目报告书-智析万卷：面向ESG报告的自适应布局分析与要素提取-专业团队ProMax.docx"
out = root / "项目报告书-智析万卷：面向ESG报告的自适应布局分析与要素提取-专业团队ProMax-国赛语言逻辑修订版.docx"

d0 = Document(src)
d1 = Document(out)


def chapter_bounds(doc):
    ps = doc.paragraphs
    s = next(i for i,p in enumerate(ps) if p.text == "ESG报告数据处理与指标体系构建")
    e = next(i for i,p in enumerate(ps) if p.text == "ESGMiner平台设计与实现")
    return s, e


def heading_tree(doc):
    s,e = chapter_bounds(doc)
    rows = []
    for p in doc.paragraphs[s:e]:
        if p.style.style_id in {"a", "a0", "gao"} or p.text.startswith("3.4.3 "):
            rows.append((p.style.style_id, p.text))
    return rows


tree0 = heading_tree(d0)
tree1 = heading_tree(d1)
s,e = chapter_bounds(d1)
text = "\n".join(p.text for p in d1.paragraphs[s:e])

print("FILE", out.name, out.stat().st_size)
print("HEADING_TREE_EQUAL", tree0 == tree1)
print("PARAGRAPHS", len(d0.paragraphs), len(d1.paragraphs))
print("TABLES", len(d0.tables), len(d1.tables))
print("SHAPES", len(d0.inline_shapes), len(d1.inline_shapes))
print("FORMULAS", [p.text.strip() for p in d1.paragraphs[s:e] if p.text.strip().startswith("（公式")])

print("\nHEADINGS")
for style, title in tree1:
    print(style, title)

required = {
    "source_manifest": "公司名称、证券代码、报告年份、公告标题、来源地址及下载时间" in text,
    "json_csv_split": "后端API返回JSON" in text and "导出中心提供CSV" in text,
    "gray_page": "灰区页面" in text and "不直接生成最终字段值" in text,
    "tfidf_fallback": "字符n-gram TF-IDF余弦检索" in text,
    "batch_budget": "默认每批处理6个字段" in text and "最多调用大语言模型12次" in text,
    "year_boundary": "默认以2024年为目标年份" in text,
    "quality_threshold": "置信度低于0.7" in text and "证据链得分低于50" in text,
    "metric_math": all(x in text for x in ["56项自动正确", "12项结果正确", "12项字段串扰", "3项数值错误", "1项年份错误"]),
    "metric_names": "自动直接采纳率" in text and "人工确认后可用率" in text,
    "limitation": "尚未独立评价系统识别未披露状态的能力" in text,
    "deploy_boundary": "配置式部署和标准化交付" in text and "不表述为零配置" in text,
}
print("\nREQUIRED")
for k,v in required.items():
    print(k, v)

print("\nABSOLUTE_WORDS")
absolute = ["确保", "绝对", "显著", "极大", "指数级", "无遗漏", "高可信度", "完整可靠", "巨大挑战", "精准", "模型性能完整"]
for word in absolute:
    print(word, text.count(word))

print("\nPROHIBITED_CLAIMS")
for word in ["用户只需登录", "请求头伪装", "规避动态加载", "时间感知阈值", "Blake2b", "384维", "任意年度自动", "全量MinerU", "开箱即用"]:
    print(word, text.count(word))

with zipfile.ZipFile(out) as z:
    print("\nZIP_TEST", z.testzip())

if tree0 != tree1 or not all(required.values()):
    raise SystemExit("QA failed")
