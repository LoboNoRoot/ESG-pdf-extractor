from __future__ import annotations

from copy import deepcopy
from pathlib import Path
from shutil import copyfile
import os

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont


ROOT = Path(r"C:\Users\Lenovo\Desktop\省一改")
SRC = next(ROOT.glob("项目报告书*ProMax.docx"))
OUT = ROOT / "项目报告书-智析万卷：面向ESG报告的自适应布局分析与要素提取-专业团队ProMax-国赛专项修订版.docx"
ASSET_DIR = ROOT / "tmp" / "report_assets"
ASSET_DIR.mkdir(parents=True, exist_ok=True)
FLOW_PNG = ASSET_DIR / "企业端到端交付闭环图.png"


def make_flow_diagram(path: Path) -> None:
    width, height = 2200, 920
    img = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(img)
    font_path = Path(r"C:\Windows\Fonts\msyh.ttc")
    bold_path = Path(r"C:\Windows\Fonts\msyhbd.ttc")
    title_font = ImageFont.truetype(str(bold_path if bold_path.exists() else font_path), 58)
    box_font = ImageFont.truetype(str(bold_path if bold_path.exists() else font_path), 32)
    small_font = ImageFont.truetype(str(font_path), 26)

    green = "#2F7D64"
    green2 = "#4F9A78"
    mint = "#EAF5F0"
    blue = "#477FA8"
    blue_bg = "#EAF2F8"
    orange = "#D38A36"
    gray = "#5F6B66"
    line = "#8AA89B"

    title = "ESGMiner企业端到端交付闭环"
    tw = draw.textbbox((0, 0), title, font=title_font)[2]
    draw.text(((width - tw) / 2, 35), title, fill=green, font=title_font)
    draw.line((160, 125, width - 160, 125), fill=green2, width=5)

    items = [
        ("1", "批量上传PDF", "标准输入"),
        ("2", "文件校验与去重", "准入控制"),
        ("3", "异步任务处理", "状态可查询"),
        ("4", "60字段结构化输出", "证据可追溯"),
        ("5", "低可信字段复核", "人机协同"),
        ("6", "JSON接口 / CSV导出", "标准化交付"),
        ("7", "跨企业指标对比", "分析应用"),
    ]
    box_w, box_h = 430, 178
    x_positions_top = [110, 625, 1140, 1655]
    x_positions_bottom = [365, 885, 1405]
    y_top, y_bottom = 200, 475

    coords = []
    for i, item in enumerate(items):
        if i < 4:
            x, y = x_positions_top[i], y_top
        else:
            x, y = x_positions_bottom[i - 4], y_bottom
        coords.append((x, y))
        fill = mint if i < 4 else blue_bg
        outline = green if i < 4 else blue
        draw.rounded_rectangle((x, y, x + box_w, y + box_h), radius=25, fill=fill, outline=outline, width=4)
        draw.ellipse((x + 22, y + 25, x + 82, y + 85), fill=outline)
        n = item[0]
        nb = draw.textbbox((0, 0), n, font=box_font)
        draw.text((x + 52 - (nb[2] - nb[0]) / 2, y + 55 - (nb[3] - nb[1]) / 2), n, fill="white", font=box_font)
        draw.text((x + 102, y + 28), item[1], fill="#20332B", font=box_font)
        draw.text((x + 102, y + 103), item[2], fill=gray, font=small_font)

    def arrow(x1, y1, x2, y2, color=line):
        draw.line((x1, y1, x2, y2), fill=color, width=7)
        if abs(x2 - x1) >= abs(y2 - y1):
            s = 1 if x2 > x1 else -1
            draw.polygon([(x2, y2), (x2 - 22 * s, y2 - 14), (x2 - 22 * s, y2 + 14)], fill=color)
        else:
            s = 1 if y2 > y1 else -1
            draw.polygon([(x2, y2), (x2 - 14, y2 - 22 * s), (x2 + 14, y2 - 22 * s)], fill=color)

    for i in range(3):
        x, y = coords[i]
        nx, ny = coords[i + 1]
        arrow(x + box_w + 10, y + box_h / 2, nx - 10, ny + box_h / 2)
    x, y = coords[3]
    nx, ny = coords[4]
    arrow(x + box_w / 2, y + box_h + 8, x + box_w / 2, y_bottom - 25)
    draw.line((x + box_w / 2, y_bottom - 25, nx + box_w / 2, y_bottom - 25), fill=line, width=7)
    arrow(nx + box_w / 2, y_bottom - 25, nx + box_w / 2, ny - 8)
    for i in range(4, 6):
        x, y = coords[i]
        nx, ny = coords[i + 1]
        arrow(x + box_w + 10, y + box_h / 2, nx - 10, ny + box_h / 2, color=blue)

    footer_y = 755
    draw.rounded_rectangle((170, footer_y, width - 170, footer_y + 105), radius=22, fill="#FFF7EA", outline=orange, width=4)
    footer = "异常闭环：准入跳过 ｜ 解析降级 ｜ 检索回退 ｜ 预算耗尽标记 ｜ 低可信结果复核"
    fb = draw.textbbox((0, 0), footer, font=small_font)
    draw.text(((width - (fb[2] - fb[0])) / 2, footer_y + 34), footer, fill="#8C5A1E", font=small_font)
    img.save(path, quality=95, dpi=(220, 220))


def find_exact(doc: Document, text: str):
    found = [p for p in doc.paragraphs if p.text.strip() == text.strip()]
    if not found and len(text.strip()) >= 18:
        prefix = text.strip()[:18]
        found = [p for p in doc.paragraphs if p.text.strip().startswith(prefix)]
    if len(found) != 1:
        raise RuntimeError(f"Expected one paragraph {text!r}, found {len(found)}")
    return found[0]


def set_text(p, text: str, style=None):
    p.text = text
    if style is not None:
        p.style = style
    return p


def remove_paragraph(p):
    parent = p._p.getparent()
    parent.remove(p._p)


def insert_before(anchor, text: str = "", style=None):
    p = anchor.insert_paragraph_before(text)
    if style is not None:
        p.style = style
    return p


def insert_after(anchor, text: str = "", style=None):
    new_p = OxmlElement("w:p")
    anchor._p.addnext(new_p)
    from docx.text.paragraph import Paragraph
    p = Paragraph(new_p, anchor._parent)
    if style is not None:
        p.style = style
    if text:
        p.add_run(text)
    return p


def replace_caption_title(p, title: str):
    field_end_seen = False
    after_runs = []
    for run in p.runs:
        fld = run._r.xpath(".//w:fldChar/@w:fldCharType")
        if "end" in fld:
            field_end_seen = True
            continue
        if field_end_seen and run.text:
            after_runs.append(run)
    if not after_runs:
        raise RuntimeError(f"Caption title run not found: {p.text}")
    after_runs[0].text = "  " + title
    for run in after_runs[1:]:
        run.text = ""


def add_seq_caption_before(anchor, label: str, cached_number: int, title: str, style):
    p = insert_before(anchor, style=style)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run(label)
    p.add_run(" ")
    r = p.add_run()._r
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    r.append(begin)
    instr_run = p.add_run()._r
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = f" SEQ {label} \\* ARABIC "
    instr_run.append(instr)
    sep_run = p.add_run()._r
    sep = OxmlElement("w:fldChar")
    sep.set(qn("w:fldCharType"), "separate")
    sep_run.append(sep)
    p.add_run(str(cached_number))
    end_run = p.add_run()._r
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    end_run.append(end)
    p.add_run("  " + title)
    return p


def clear_picture_paragraph(p, image_path: Path):
    p.text = ""
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(str(image_path), width=Cm(15.6))


def shade_cell(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_text(cell, text: str, bold=False, color=None, size=8.5, align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(text)
    run.bold = bold
    run.font.name = "微软雅黑"
    run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor(*color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


if OUT.exists():
    os.chmod(OUT, 0o666)
copyfile(SRC, OUT)
os.chmod(OUT, 0o666)
make_flow_diagram(FLOW_PNG)
doc = Document(OUT)

body_style = find_exact(doc, "ESG报告通常以PDF格式进行公开披露，非结构化的文档形态使得信息提取面临巨大挑战。在遵循合法合规原则的基础上，本方案的数据采集工作主要以中国证监会指定的信息披露平台巨潮资讯网（www.cninfo.com.cn）作为核心数据来源，通过网络爬虫采集到200份数据，涵盖金融、房地产、制造、能源、科技、消费等多个行业。每份报告均按照统一规则进行命名和存储，并同步记录公司名称、股票代码、报告年份、文件来源等基础信息，便于后续数据管理和来源追踪。").style
heading3_style = find_exact(doc, "过滤与扫描模块").style
caption_style = find_exact(doc, "图 2  ESG指标体系框架图").style
chapter3_body_style = find_exact(doc, "在面对中文ESG专业术语时，纯词汇匹配方法难以捕捉专业领域、相似含义词汇语义表征之间的等价关系，而通用嵌入模型对专业领域的细粒度区分又往往不够精准。为此，本方案设计了一种无监督的查询改写驱动双路融合检索架构，在不依赖外部标注数据的前提下，同时兼顾精准词汇匹配与深层语义理解。").style

# ---------------- Chapter 2 ----------------
ch2_21 = find_exact(doc, "数据获取")
p = ch2_21._p.getnext()

ch2_data_old = [
    "ESG报告通常以PDF格式进行公开披露，非结构化的文档形态使得信息提取面临巨大挑战。在遵循合法合规原则的基础上，本方案的数据采集工作主要以中国证监会指定的信息披露平台巨潮资讯网（www.cninfo.com.cn）作为核心数据来源，通过网络爬虫采集到200份数据，涵盖金融、房地产、制造、能源、科技、消费等多个行业。每份报告均按照统一规则进行命名和存储，并同步记录公司名称、股票代码、报告年份、文件来源等基础信息，便于后续数据管理和来源追踪。",
    "通过模拟浏览器行为向服务器发送POST请求，使用常用的爬虫工具，批量抓取下载符合关键词“ESG报告”“环境、社会和公司治理报告”“社会责任报告”“可持续发展报告”的PDF文档。针对巨潮资讯网的动态加载特性，本方案的技术采用基于urllib库的底层HTTP通信机制，模拟常见的用户代理请求头，并在请求之间设置随机延时，以减少高频访问带来的访问限制风险。程序首先调用公告查询接口分页获取公告元数据，再根据公告ID与附件地址拼接PDF下载链接，将文件保存到本地数据目录。对于下载失败、请求超时或文件损坏的情况，程序会记录失败日志，并通过重试机制再次下载。",
    "在本方案架构中，上传的PDF首先会经过一个文档准入过滤层级，采用设定关键词词典来排除不相关文档并提醒用户。被纳入的文档则会进入下一层文档扫描环节，这样的机制有利于提高提取效率，也避免了大量空值对数据集的污染。",
]
ch2_data_new = [
    "本方案以公开披露的ESG相关报告为处理对象，核心数据来源为中国证监会指定信息披露平台巨潮资讯网（www.cninfo.com.cn）。采集范围覆盖金融、房地产、制造、能源、科技、消费等行业，共形成200份报告样本。采集模块按照“公告检索—元数据登记—PDF下载—本地落盘”的顺序运行，并为每份报告建立来源清单，至少记录公司名称、证券代码、报告年份、公告标题、来源地址和下载时间，为后续复核、更新与责任追溯提供依据。",
    "采集模块按照公开披露接口规范设置必要请求头，通过访问间隔、失败重试和状态记录维持批量任务的连续性。文件落盘后，系统检查扩展名、文件大小、PDF可读性及下载完整性；请求超时、下载失败或文件损坏时记录原因并进入重试或人工补采队列。该过程将“已发现公告”与“可进入处理的数据文件”分开管理，避免异常文件直接进入抽取环节。",
    "进入ESGMiner平台前，文档准入模块还会依据文件名与前两页文本进行目标报告识别，并对重复文件执行SHA-256校验。非目标报告生成结构化跳过记录，重复文件返回既有任务信息，通过校验的PDF才进入异步处理队列。该环节属于产品输入质量控制，具体判定规则与页面处理方法将在第三章展开。",
]
for old, new in zip(ch2_data_old, ch2_data_new):
    set_text(find_exact(doc, old), new, body_style)

ch2_schema_old = [
    "本方案构建了一套面向A股上市公司的跨行业核心ESG指标体系，共计60个关键字段，严格遵循高频披露、可结构化抽取、可人工复核、跨行业可比的设计原则。该体系不追求穷尽所有行业细节，而是聚焦监管关注度高、报告中出现频率稳定、具备明确证据支撑的核心指标，形成稳定且可扩展的量化分析基础。",
    "指标体系按环境、社会、治理三大维度搭建，覆盖气候变化、资源利用、污染防治、员工权益、供应链管理、创新驱动、公司治理、商业道德等核心议题。其中环境维度22项，重点围绕温室气体排放、能源消耗、水资源管理、废弃物处理及生态保护展开；社会维度28项，聚焦员工结构与发展、安全生产、职业健康、供应链责任、客户与社区关系；治理维度10项，覆盖董事会结构、廉洁合规、风险管理、信息披露与创新管理等内容，如图 2所示。",
    "为适应报告中披露形式的多样性，本方案将指标划分为23项定量指标、27项定性指标与10项混合指标（图 2，图 3）。定量指标侧重数值、单位与年份的精准映射，例如温室气体排放总量、员工总数、研发投入等；定性指标侧重治理机制、政策措施的语义提炼，例如气候变化治理机制、员工健康与安全管理等；混合指标则兼顾指标数值与配套机制，例如客户投诉解决、信息安全事件等。该分类为后续差异化抽取策略提供了明确依据。",
    "每个指标均配备完整元数据画像，包括中英文名称、所属议题、别名集合、检索关键词、领域知识、期望单位、必备证据约束及排除词典。这些画像共同构成字段配置，使模型在RAG检索阶段实现打分排序，提升证据召回的针对性与准确性。同时，通过代码中的规则强约束，有效抑制字段间混淆，解决同表格中相邻行字段相互干扰的痛点。",
    "指标选择上，本方案采用优先纳入A股ESG报告高频披露项，兼顾定量绩效与定性机制，确保跨公司可比性；其次评估技术可抽取性，确保在后续解析提取流程中具备稳定证据链；最终保留人工复核友好性，支持字段级溯源与修改。",
    "通过这一核心指标体系，本方案实现了从非结构化PDF到结构化JSON或CSV结果的端到端映射，为后续投资决策、监管审查与学术分析提供了高可信度的量化基础。",
]
ch2_schema_new = [
    "本方案面向A股上市公司构建60项跨行业核心ESG指标，指标选择依次遵循高频披露、可结构化表达、可人工复核和跨企业可比四项原则。体系不追求穷尽行业全部披露事项，而是优先纳入监管与市场关注度较高、报告中具有可定位证据、能够形成统一字段口径的指标，作为企业批量治理与持续维护的数据标准。",
    "指标体系由环境、社会、治理三类组成：环境维度22项，覆盖温室气体排放、能源、水资源、废弃物及生态保护；社会维度28项，覆盖员工发展、安全健康、供应链、客户与社区责任；治理维度10项，覆盖董事会结构、廉洁合规、风险管理与信息披露。三类指标及其议题关系如图 2所示。",
    "按照披露形态，60项指标进一步分为23项定量指标、27项定性指标和10项混合指标（图 2、图 3）。定量指标要求数值、单位与年份对应；定性指标记录机制、政策或举措的证据摘要；混合指标同时保留量化结果与配套说明。分类结果用于约束字段输出类型，并使不同企业报告能够在同一口径下进入复核和比较流程。",
    "每个字段均由可配置Schema描述，元数据包括标准名称、别名、字段类型、期望单位、目标年份、检索关键词、证据约束和排除词。字段配置将检索线索、抽取要求与质量判断统一在同一数据字典中，便于企业根据披露口径维护别名与单位规则，也便于同一配置复用于不同报告和后续年度批次。",
    "指标的结构化结果同时保留字段值、单位、年份、来源页码、证据文本、置信度与复核状态。后端API以JSON返回任务和字段结果，导出中心面向业务用户生成CSV文件；二者分别服务系统集成与离线分析，避免接口返回与文件下载口径混用。",
    "由此形成的指标体系不仅是抽取清单，也是连接输入文档、证据召回、质量复核和企业数据应用的字段标准，为跨报告统一口径和结果追溯提供可维护的数据基础。",
]
for old, new in zip(ch2_schema_old, ch2_schema_new):
    set_text(find_exact(doc, old), new, body_style)

set_text(find_exact(doc, "本方案针对ESG报告篇幅冗长、信息密度极不均衡的特点，设计了一套以降噪和分流为核心的自适应布局分析与要素提取架构。该架构通过文档级过滤、页面级评估、混合分流解析、LLM灰区仲裁以及字段级RAG抽取五个流程层次，构建了从规则驱动到模型辅助的完整决策链路。模型在流程中不直接处理整份PDF，而是逐级缩小高成本处理范围，以达到资源投入与提取质量之间的动态平衡。"), "ESGMiner平台采用分层处理方案，将长篇、复杂版式PDF拆分为输入控制、证据构建、结构化抽取和交付复核四个环节。第二章仅说明各环节的业务目的、输入与输出；页面评分、检索排序和质量阈值等具体实现留至第三章展开。", body_style)

set_text(find_exact(doc, "过滤与扫描模块"), "输入与准入控制", heading3_style)
set_text(find_exact(doc, "模型先对上传PDF执行轻量级准入验证。读取文件名及前两页文本，去除空白后比对ESG正向词集与排除词，仅当同时满足“含ESG框架信号”且“不含非目标特征词”时方可进入后续流程。该机制通过前置过滤有效拦截异常文档，既节约解析时间，又避免无效抽取产生空值污染数据集，并对跳过文档生成包含原因代码、文件路径与状态标识的结构化审计记录。"), "输入为企业批量上传的PDF及其文件元数据。系统依次完成格式与大小校验、PDF可读性检查、SHA-256重复识别和目标报告准入判断；通过校验的文件生成独立任务编号并进入异步队列，未通过的文件保留跳过或冲突原因，用户可在任务列表查看处理状态。", body_style)
set_text(find_exact(doc, "经过初步过滤后，符合规则的ESG文档将进入页面价值评估层级。ESG报告的结构决定了全量高精度解析将产生算力浪费，封面、目录、宣传致辞等低价值页面占比极高，而承载核心量化指标的表格页与图文混排页仅为少数。若对所有页面一视同仁地投高资源视觉版面分析模型，将导致计算成本指数级增长。为此，本方案采用规则评分进行全文扫描，对页面信息密度进行量化估计。"), "准入环节的输出不是字段结果，而是可处理任务、文件指纹和异常状态。这样既将无关文件隔离在高成本处理之外，也为批量重试、重复识别和运行审计提供统一入口。", body_style)

# Remove Chapter 2's implementation-detail table and its caption.
old_table1_caption = find_exact(doc, "表 1  三大核心维度")
remove_paragraph(old_table1_caption)
old_table1 = doc.tables[0]
old_table1._tbl.getparent().remove(old_table1._tbl)
for t in [
    "基于上述逻辑，本方案采用快速从文档中系统性提取三大核心维度共7个特征，并进行量化。如表 1：",
    "实践中，高分页面往往对应ESG数据汇总表、关键绩效披露页或低质量扫描页；低分页面则多为封面、目录、宣传叙述段落。这一轻量级预筛选使模型在无需调用任何大模型的前提下，即可完成全文页面的价值估计，为后续混合分流解析策略提供决策依据。",
]:
    remove_paragraph(find_exact(doc, t))

set_text(find_exact(doc, "选择性解析模块"), "分层解析与证据构建", heading3_style)
set_text(find_exact(doc, "在此层级，即以极低的计算代价在流程早期阶段筛去无价值或低价值页面，仅对高密度信息承载页投入昂贵资源，形成从文档级到页面级、从规则驱动到模型辅助的递进式决策链路。模型构建了多层分流机制，实现页面解析策略的动态自适应分流。具体而言，分流决策基于页面信息密度综合评估，当页面规则得分高于上阈值时，直接判定为高价值页面；低于下阈值时，进入轻量级通道；处于两阈值之间的灰区页面由视觉LLM进行二次语义判断，从而降低单一规则误判。"), "分层解析以页面路由结果为输入。系统保留全文基础文本，并将高价值或复杂版式页面选择性送入深度解析；灰区页面的文本或视觉模型仅判断该页是否需要进入深度解析通道，不直接生成最终ESG字段结果。两路解析结果最终统一为带来源、原页码和块编号的证据文本块。", body_style)
set_text(find_exact(doc, "文档中，低价值页面就由轻量快速的PyMuPDF库来提取，利用PDF原生文本层进行高效提取，吞吐量优势显著，适用于封面、目录及纯文本叙述页。复杂页面交由高精度文档识别工具MinerU，该工具的架构能够准确识别复杂多列布局、合并单元格表格、图文混排区域及跨页连续结构，并输出保留行列关系与层次结构的Markdown格式文本。在处理扫描图像页或物理版面顺序混乱的报告时，MinerU的多模态信息提取能力可有效还原表格拓扑与视觉上下文，确保后续字段抽取的证据完整性。两个分支的最终结果会经过统一规范化处理，随后进入文本切块，进行融合。"), "采用分层方案的原因是报告各页的信息价值与版式复杂度差异较大。基础通道提供全文可检索文本，深度通道补充表格、图文混排与扫描候选页的结构信息；当深度解析未启用或运行失败时，任务仍保留基础文本结果和降级状态，便于继续处理或人工判断。", body_style)
set_text(find_exact(doc, "对于难以判定的灰区文档，则会交由视觉LLM模型进行精细解析提取，输出的JSON数据经过语义连贯性保持的自适应切块后，直接服务于指标构建与模型决策环节，无需输入前两个分支中重复解析或额外切块。"), "该环节向下游提供统一证据库，而不提前做字段结论。证据块的来源、页码和块编号将随抽取结果返回，使企业复核人员能够从结构化字段回到报告原文。", body_style)

set_text(find_exact(doc, "基于混合检索的字段抽取模块"), "结构化抽取与质量闭环", heading3_style)
set_text(find_exact(doc, "混合解析结果汇合后，模型构建RAG标准化文本块库。随后，基于预定义的60字段ESG指标体系，执行字段级结构化抽取。"), "结构化抽取以60项字段Schema和统一证据库为输入。系统为每个字段检索候选证据，再由大语言模型按批次输出字段值、单位、年份、页码、证据和置信度；未披露或证据不足时返回未命中状态，不以推测内容补齐字段。", body_style)
set_text(find_exact(doc, "文本切分与索引构建采用滑动窗口机制，确保上下文连续性。模型在异构文本块上构建紧凑索引，利用改进BM25算法融合字段中英文名、别名、领域知识及预期单位等特征，实现高效证据召回。在Top-K证据选择时，引入硬编码奖励机制，当文本块包含特定数据表头、计量单位连续短语或必备领域词时给予显著加分，确保高价值表格区域优先进入抽取上下文，同时对包含排除词的文本块进行降权，以降低字段串扰。"), "抽取后系统执行年份、单位、数值格式、来源页码和证据完整性检查，并将置信度或证据链不足的字段标记为待复核。用户可对字段值及其证据进行确认或修改，复核结果与原始输出同时保留，形成“自动抽取—质量标记—人工确认—结果交付”的闭环。", body_style)
set_text(find_exact(doc, "考虑到实际应用需求，本方案将语义类别权重与时间因素注入向量空间，使RAG检索结果自然倾向近期关键ESG事项。针对这一问题，本方案引入了时间感知阈值机制，可根据查询时间范围自动调节余弦相似度阈值，有效解决历史数据与近期数据的检索平衡问题。"), "API面向系统集成返回JSON结果，导出中心面向业务使用生成CSV。两类输出均携带复核状态，企业可将自动结果与人工确认结果分开使用，避免低可信字段在未核验情况下进入分析。", body_style)
for t in [
    "本方案所有RAG相关代码均实施统一依赖注入，确保各模块间保持松耦合关系，提升RAG系统可维护性。为进一步保障数据质量，模型集成后处理校验机制。该机制包括单位归一化映射，支持别名转换与纠偏、数值正则提取，实际年份与设定年份对齐验证以及基于名称、数值、单位、页码等维度的证据链评分模型。证据链评分低于50分的字段自动标记为弱证据链，提示用户优先进行人工复核。",
    "通过RAG字段抽取算法，本方案实现了从非结构化PDF到结构化JSON/CSV结果的端到端提取，输出结果具备完整溯源性，显示证据原文、页码与来源路径，为投资者或机构分析使用、监管审查及学术研究提供了高可信度的ESG量化基础数据。",
]:
    remove_paragraph(find_exact(doc, t))

# Replace and move Figure 4 to the enterprise delivery subsection.
fig4_caption = find_exact(doc, "图 4  分流决策核心分支图")
fig4_image_p = fig4_caption._p.getprevious()
from docx.text.paragraph import Paragraph
fig4_image = Paragraph(fig4_image_p, fig4_caption._parent)
clear_picture_paragraph(fig4_image, FLOW_PNG)
replace_caption_title(fig4_caption, "企业端到端交付闭环图")

ch2_eval = find_exact(doc, "性能评估指标")
insert_before(ch2_eval, "企业化交付流程", heading3_style)
insert_before(ch2_eval, "企业完成运行环境和模型服务参数配置后，可通过标准化部署方式启动系统，并直接使用报告上传、任务处理、字段复核、结果导出和指标分析功能。该定义强调配置式部署与快速上线，不包含用户认证、任意模型零配置接入等现阶段未实现能力。", body_style)
insert_before(ch2_eval, "面向业务用户，ESGMiner平台形成“批量上传PDF—文件校验与去重—异步处理—60字段输出—低可信字段复核—CSV导出或API调用—跨企业对比”的端到端流程，如图 4所示。每一步均保留任务状态或证据记录，异常文件、降级任务和待复核字段均可被识别，而不是与正常结果混合交付。", body_style)
fig4_image._p.getparent().remove(fig4_image._p)
fig4_caption._p.getparent().remove(fig4_caption._p)
ch2_eval._p.addprevious(fig4_image._p)
ch2_eval._p.addprevious(fig4_caption._p)
insert_before(ch2_eval, "本章完成了数据来源、指标标准、产品处理流程及评估口径的总体设计。下一章将在此基础上，进一步说明文档过滤、重点页解析、证据召回、大模型抽取与质量校验的具体实现，并结合人工核验结果评价系统性能。", body_style)

# Chapter 2 evaluation section.
set_text(find_exact(doc, "评估以报告与指标的交叉矩阵为基本单位。由于部分企业在报告中确实未对某些指标进行披露，此类情况不应纳入抽取精度评估，否则会导致评估结果虚低。为此，本方案对“报告未披露”项目单独标注，并从总核验项中剔除，剩余部分定义为有效样本："), "评估以“报告—指标”交叉项为基本单位。人工核验先判断原报告是否披露该字段：已披露项进入字段值准确性评价；“未披露”作为独立状态单独记录，不进入已披露字段的值准确率分母。有效样本定义如下：", body_style)
set_text(find_exact(doc, "其中，核验报告数量为，核验指标数量为，则理论上总核验项为。"), "式（1）中，N_total表示总核验项，N_report表示核验报告数量，N_field表示每份报告的核验指标数量，N_undisclosed表示经人工确认未披露的项目数，N_effective表示已披露有效样本数。", body_style)
set_text(find_exact(doc, "后续，本方案的全部性能指标均以有效样本为分母计算，以保证评估口径的合理性与可比性。"), "抽取值相关指标以N_effective为分母；同时单列未披露项目的系统识别结果。该口径适合评价已披露字段的抽取质量，但不能据此评价系统识别“未披露”的能力，因此第三章将明确披露这一评估边界。", body_style)
set_text(find_exact(doc, "抽取准确率衡量模型对有效披露指标的命中质量，具体定义为："), "自动直接采纳率用于衡量系统输出中无需人工修改、可直接采纳的正确结果占有效样本的比例，沿用式（2）的计算形式：", body_style)
set_text(find_exact(doc, "其中，正确数量的判定标准为字段值、计量单位与所属年份均与原报告一致。"), "式（2）中，N_direct表示字段值、计量单位、所属年份与原报告一致，且未被质量机制标记为待复核的结果数量。对于结果正确但被标记待复核的项目，另计入“人工确认后可用率”，不与自动直接采纳率混称。", body_style)
set_text(find_exact(doc, "漏抽率衡量报告中明确存在、但模型未能召回的指标占比："), "漏抽率衡量报告中明确披露、但系统未返回有效候选结果的项目占比，计算见式（3）。漏抽记录应与未披露状态分开，避免将“原文没有”误作“系统未找到”。", body_style)
set_text(find_exact(doc, "字段串错率是本方案重点关注的专项指标，用于度量模型将相邻或语义相近指标的证据错误映射到目标字段问题，反映检索与抽取阶段的字段边界区分能力："), "字段串扰率用于度量系统将相邻或语义相近指标的证据错误映射到目标字段的情况，反映检索与抽取阶段的字段边界区分能力，计算见式（4）：", body_style)
set_text(find_exact(doc, "该指标在ESG报告场景中具有较强的诊断意义，因为“范围一排放”与“范围二排放”、“员工总数”与“女性员工占比”等相邻指标在原文中往往位于同一表格的相邻行列，检索层若召回边界不清晰，极易引发字段串扰，进而影响下游抽取的准确性。"), "除字段串扰外，评估还分别记录数值错误与年份错误：前者指字段归属正确但数值或单位不一致，后者指字段值对应的统计年度不符合目标年份。错误类型采用互斥标注，使各类数量能够汇总到有效样本总数。", body_style)
set_text(find_exact(doc, "本方案以抽取准确率为主要优化目标，以字段串错率为重点监控指标（表 2），评估流程全程保留每条样本的人工标注记录、原始抽取结果与错误类型分类，确保评估结果具备完整的可追溯性，支持后续迭代优化的误差定位与分析。"), "评价口径汇总见表 1。评估过程保留人工标注、原始输出、复核状态与互斥错误类型，分别报告自动直接采纳率、人工确认后可用率、漏抽情况、字段串扰率、数值错误和年份错误，便于定位系统改进方向。", body_style)

# Existing performance table is now Table 1 after removing the old implementation table.
perf_table = doc.tables[0]
headers = ["性能指标", "计算口径", "说明"]
rows = [
    ("自动直接采纳率", "直接正确结果/有效样本", "无需人工确认即可采纳"),
    ("人工确认后可用率", "直接正确与复核确认正确/有效样本", "体现人机协同后的可用结果"),
    ("漏抽率", "漏抽/有效样本", "已披露但未返回有效候选"),
    ("字段串扰率", "字段串扰/有效样本", "相邻或近义字段边界混淆"),
    ("数值错误", "互斥错误计数", "值或单位与原报告不一致"),
    ("年份错误", "互斥错误计数", "结果不属于目标统计年度"),
    ("未披露识别", "独立记录，不计入值准确率", "需另行评价未披露判断能力"),
]
while len(perf_table.rows) < len(rows) + 1:
    perf_table.add_row()
for j, h in enumerate(headers):
    set_cell_text(perf_table.cell(0, j), h, bold=True, color=(255, 255, 255), align=WD_ALIGN_PARAGRAPH.CENTER)
    shade_cell(perf_table.cell(0, j), "2F7D64")
for i, row in enumerate(rows, 1):
    for j, value in enumerate(row):
        set_cell_text(perf_table.cell(i, j), value, align=WD_ALIGN_PARAGRAPH.CENTER if j == 0 else WD_ALIGN_PARAGRAPH.LEFT)
        shade_cell(perf_table.cell(i, j), "EEF6F2" if i % 2 else "FFFFFF")
replace_caption_title(find_exact(doc, "表 2  性能指标说明"), "性能评估口径说明")

# ---------------- Chapter 3 ----------------
ch3 = find_exact(doc, "基于RAG与大模型的ESG指标结构化提取算法")
ch3_first_h2 = find_exact(doc, "报告筛选与页面价值评估")
insert_before(ch3_first_h2, "第二章界定了数据来源、60项指标标准、产品处理流程与评估口径。本章进入实现层，按照“报告筛选与页面路由—分层解析—字段级证据召回—批次化抽取与质量校验—人工核验”的顺序说明各模块的输入、处理、输出及异常路径。", chapter3_body_style)

set_text(find_exact(doc, "文档准入过滤模块以PDF文件名与首页文本为输入，执行基于规则的二元分类判断。模块设置了两个关键词集合，即ESG标题正向词集与排除词集，并终止后续全部处理流程。此机制有效防止异常输入污染评估集与抽取准确率统计。"), "报告有效性过滤以PDF文件名和前两页文本为输入。系统规范空白字符后，分别匹配ESG标题正向词集与排除词集：命中正向信号且未命中排除条件时进入页面扫描，否则生成跳过记录并终止该文件的后续处理。", chapter3_body_style)
set_text(find_exact(doc, "在正式解析之前，模型首先对上传的PDF文档执行一轮轻量级的验证，判断其是否属于可处理的ESG报告范畴。模型读取文件名与PDF前两页的文本内容，在去除多余空白字符后，与两个预设的关键词集合逐一比对。一是ESG标题正向词集，包含ESG、可持续发展、环境、社会及治理、社会责任等表述；二是排除词集，用于识别鉴证声明、摘要、临时公告等非目标文档。只有当文档命中正向词且未命中排除词时，才进入后续页面扫描流程。"), "正向词用于识别ESG、可持续发展、环境、社会及治理、社会责任等目标标题；排除词用于识别鉴证声明、摘要、临时公告等非目标文件。该规则由系统模块执行，输出为“通过”或带原因代码的“跳过”状态，不调用大语言模型。", chapter3_body_style)
set_text(find_exact(doc, "对于未通过准入的文档，模型会生成一份结构化的跳过记录JSON文件，注明具体的跳过原因代码，并立即终止后续全部处理流程，避免将无效输入引入高成本的解析与大模型调用环节。"), "未通过准入的文件保留结构化原因和任务状态，便于用户查看或重新提交；通过准入的文件才进入页级特征扫描。由此将输入质量问题与后续抽取错误分开记录。", chapter3_body_style)

set_text(find_exact(doc, "ESG报告中真正承载量化指标的页面，往往具有几个共同的文本特征，比如数字密集、含有关键绩效、单位等表格信号词、出现碳排放、员工、治理等ESG专业术语，以及大量年份标注。与此同时，扫描版或图片页因文本层几乎为空，同样需要被标记出来以便后续专门处理。反之，目录、封面和宣传致辞等页面通常文本结构简单、指标密度较低。基于这些观察，本方案设计了一套轻量级的页级评分机制，用于量化每一页的信息价值。"), "系统对通过准入的报告逐页提取七类特征：文本长度、数字出现频次、年份命中、ESG术语命中、表格信号词命中、类表格行特征和目录索引惩罚。数字、年份和表格信号用于刻画指标密度；低文本页仅作为待进一步判断的候选页，不能直接推断为高价值扫描页。", chapter3_body_style)
set_text(find_exact(doc, "各项得分上限经过校准，以防单一特征主导结果，最终得分截断为非负值，得分越高，代表该页承载有效ESG数据的可能性越大。具体的评分规则与阈值如表 3所示："), "各项得分设置上限并截断为非负值，得分越高表示该页承载ESG证据的可能性越大。评分规则与当前阈值见表 2，七类特征的加权关系见式（5）。", chapter3_body_style)
set_text(find_exact(doc, "该评分用于后续页面分流决策。实践中，高分页面往往对应ESG数据汇总表、关键绩效披露页或低质量扫描页；低分页面则多为封面、目录、宣传叙述段落。基于上述逻辑，模型对相应页面提取对应三个维度的七类文本特征，并加权求和得到信息密度评分："), "评分结果仅用于页面路由。高分页通常对应数据汇总表或关键绩效披露页；低分页通常对应封面、目录或叙述页；低文本页和分数居中的页面仍需进入灰区判断。页面信息密度得分定义为：", chapter3_body_style)
set_text(find_exact(doc, "其中，为第p页的分数；为示性函数，通过0-1二值逻辑来标识变量；表示文本极少，不足规定值则取1；为表格信号词命中数；为ESG关键词命中数，表示ESG关键词大于等于3，满足取1；表示数字密集程度；为类表格行数；为年份关键词命中数量；是目录页惩罚标识，存在索引且数字较多的页面判定为目录并扣分。"), "式（5）中，S_p表示第p页得分；I(·)为示性函数；L_p、T_p、E_p、D_p、R_p、Y_p分别表示低文本标识、表格信号词命中、ESG关键词命中、数字密集程度、类表格行数和年份命中数量；C_p为目录页惩罚标识。", chapter3_body_style)
set_text(find_exact(doc, "具体而言，模型对每页文本执行快速特征识别扫描，提取七类关键指标，包括文本长度、数字出现频次、年份标注命中、ESG专业术语命中、表格信号词命中、类表格行特征以及目录索引惩罚标识。这些特征共同构成了页面信息价值的多元表征，从数字密度、语义相关性与版面视觉拓扑三个核心维度反映不同页面的结构化信息潜力。"), "系统将各页得分、特征摘要和原页码写入路由结果。该结果既可供后续选择重点页，也可用于复核为何某页进入深度解析通道。", chapter3_body_style)
set_text(find_exact(doc, "本方案的这一设计在实际工程落地中具有显著的降本增效价值。一方面，它以纯规则驱动的计算方式替代了全量高精度视觉解析，将算力消耗压缩至可接受范围，使模型具备处理数百页级长文档的规模化能力；另一方面，它为后续混合解析策略的动态分流奠定了数据基础，确保有限的深度解析资源能够优先作用于最具数据价值的页面。"), "采用规则先扫描、重点页再深度解析，可将资源优先投入信息密度较高的页面。深度解析成本会随页数持续累积，因此系统还设置重点页数量上限，以防单份长报告占用过多资源。", chapter3_body_style)

set_text(find_exact(doc, "规则评分能够清晰地识别数据表密集的高价值页面和封面、致辞等低价值页面，但对于得分处于中间区间的页面，单靠数字指标往往难以准确判断，本方案将这类页面定义为灰区页面。例如，一页图文并排的环境绩效说明，可能因为图片占据大量版面导致文本稀疏、得分偏低，但实际上包含重要信息；反之，一页罗列大量年份的目录页可能因年份密集得分偏高，却不具备实际抽取价值。"), "规则评分对高分和低分页面给出初始路由，但中间区间以及部分低文本候选页可能存在歧义，本方案统一称为灰区页面。例如，图文并排的环境绩效页可能因文本稀疏而得分偏低，年份密集的目录页也可能获得较高分数。", chapter3_body_style)
set_text(find_exact(doc, "为此，本方案引入一套分流机制，如图 5所示，以阈值和将全文页面划分为三个区间，并分别采取不同的处理策略。如表 4，对于落入灰区的页面，模型将其文本长度、数字数量、ESG关键词命中、表格信号词命中等特征摘要打包成结构化提示，交由一个轻量级语言模型进行语义层面的二次判定，仅返回“是/否”与置信度，若模型不可用则回退至规则得分结果。"), "系统按上下阈值将页面划分为高价值区、灰区和低价值区，分流机制见图 5、表 3。灰区复判向语言模型或视觉模型提供页面特征摘要，并只接收“是否进入MinerU”及置信度等路由信息；模型不可用、未启用或超过复判页数上限时，系统回退到规则结果并记录回退状态。", chapter3_body_style)
set_text(find_exact(doc, "这一设计的关键优势在于分层介入、按需调用。规则层以极低的计算成本完成绝大多数页面的分流，LLM仅对规则层无法确定的少数模糊页面介入裁决，从而将大模型的调用总量严格控制在设定的上限之内。"), "灰区模型不输出最终ESG字段，也不替代第三节的证据召回。其唯一职责是补充页面路由决策，使深度解析调用保持在可配置范围内。", chapter3_body_style)

# 3.2 implementation order.
set_text(find_exact(doc, "页面分流与解析策略"), "页面路由结果生成", heading3_style)
set_text(find_exact(doc, "完成三档分流决策后，模型进入实际的解析执行阶段。全文页面被最终归入两条处理通道，二者互不重叠、覆盖全文所有页面。"), "页面评分与灰区复判完成后，系统输出按原页码排列的路由清单，记录每页得分、处理通道和路由依据。高价值页面优先进入深度解析候选集；当候选页超过配置上限时，仅保留排序靠前的页面，避免资源失控。", chapter3_body_style)
set_text(find_exact(doc, "MinerU解析通道负责处理被判定为高价值的页面[19]。MinerU是一款专为学术与专业文档设计的版面分析工具，能够准确识别复杂的多列布局、合并单元格表格、图文混排区域，并将其转换为结构完整的Markdown格式输出，保留表格的行列关系与文本的层次结构。考虑到ESG报告通常超过100页，若对全文执行MinerU解析，计算成本将呈指数级增长，MinerU的深度处理范围被严格限定在经页级评分与灰区仲裁确认的高价值页集合。"), "“覆盖全文所有页面”是指PyMuPDF对全文提供基础文本处理，并不表示所有页面都进入MinerU。深度解析候选页当前设置最大页数上限，用于约束复杂页面处理范围；超过上限的页面仍保留基础文本和路由记录。", chapter3_body_style)
set_text(find_exact(doc, "重点页子PDF生成"), "重点页子PDF生成", heading3_style)
set_text(find_exact(doc, "完成页面分流后，模型不会直接将整份ESG报告输入MinerU进行解析，而是先根据页面价值评估结果生成重点页列表。重点页常包括ESG指标表、环境绩效表、员工数据表、治理结构表及文本层质量较差但可能包含关键信息的页面。"), "系统依据路由清单抽取重点页，候选对象包括ESG指标表、环境绩效表、员工数据表、治理结构表以及可能包含信息的低文本页。页面选择保留原始页码，不改变证据定位口径。", chapter3_body_style)
set_text(find_exact(doc, "模型利用PyMuPDF从原始PDF中抽取这些重点页面，并按照原有页码顺序重新生成一份子PDF，该文档只包含后续需要深度解析的页面来作为MinerU的实际输入文件。通过这种方式，模型既保留了原报告中的页面版式和表格结构，又减少了MinerU的解析页数，避免对封面、目录和大段叙述性页面进行重复处理。"), "系统使用PyMuPDF按原页码顺序生成重点页子PDF，作为MinerU的输入；与此同时，原始PDF仍由PyMuPDF提取全文基础文本。子PDF只缩小深度解析范围，不替代原报告，也不改变原始页码与证据页码的映射。", chapter3_body_style)
set_text(find_exact(doc, "该步骤不仅可以降低整份报告解析带来的时间成本、计算资源开销，提高复杂页面解析的针对性，还能保留重点页与原报告页码之间的对应关系，便于后续结果追溯和人工复核。"), "子PDF生成失败时，任务记录异常并继续使用全文基础文本；生成成功后才进入MinerU，便于将文件处理错误与解析工具错误分别定位。", chapter3_body_style)
set_text(find_exact(doc, "MinerU深度解析与文本融合"), "深度解析、文本融合与降级", heading3_style)
set_text(find_exact(doc, "完成子PDF的生成后，模型调用MinerU对其中的重点页面进行深度解析。MinerU相比普通PDF文本提取工具，更适合处理ESG报告中常见的复杂版式内容，例如多列表格、图文混排、跨行跨列指标表以及扫描质量较差的页面。解析完成后，模型会读取MinerU输出的文本与结构信息，并将其重新映射至原始报告页码。"), "完整生产模式下，系统对重点页子PDF调用MinerU，补充多列表格、图文混排、跨行跨列指标表等复杂版式信息，并将结果映射回原始页码。为适应低资源环境，MinerU采用选择性调用且可配置启停，而非所有部署环境默认全量运行。", chapter3_body_style)
set_text(find_exact(doc, "由于不同解析工具各有优势，本方案没有只依赖MinerU的结果，而是将MinerU解析文本与PyMuPDF提取文本进行融合。PyMuPDF负责提供稳定的页级文本和视觉顺序文本，MinerU负责补充表格、图片化页面和复杂版式中的细节内容。融合时，模型会保留文本来源、原始页码及解析工具标识，避免两路结果混淆。"), "系统将PyMuPDF全文基础文本与MinerU重点页结果统一转换为文本块，每个文本块记录source、page、chunk_id和text。两路内容允许在同一页并存，检索阶段可根据来源和字段特征选择证据，复核阶段也能识别证据来自基础解析还是深度解析。", chapter3_body_style)
set_text(find_exact(doc, "通过该方式，后续字段检索阶段既可以利用PyMuPDF的快速文本结果，也可以利用MinerU对复杂页面的深度解析结果，从而提高ESG指标证据召回的完整性和准确性。"), "在低配演示模式、MinerU未启用或调用失败时，系统保留PyMuPDF文本块并标注降级状态，后续检索和抽取仍可继续；企业可依据任务日志决定是否补充深度解析后重跑。", chapter3_body_style)

# 3.3 retrieval.
set_text(find_exact(doc, "在面对中文ESG专业术语时，纯词汇匹配方法难以捕捉专业领域、相似含义词汇语义表征之间的等价关系，而通用嵌入模型对专业领域的细粒度区分又往往不够精准。为此，本方案设计了一种无监督的查询改写驱动双路融合检索架构，在不依赖外部标注数据的前提下，同时兼顾精准词汇匹配与深层语义理解。"), "字段级证据召回按照“文本块构建—查询改写—BM25与Embedding双路召回—RRF融合—字段重排—Top-K输出”执行。该顺序既保留专业术语的精确匹配，也补充同义表达的语义召回，并通过字段级约束减少同表相邻指标之间的证据串扰。", chapter3_body_style)
set_text(find_exact(doc, "报告文本块构建"), "文本块构建与查询改写", heading3_style)
set_text(find_exact(doc, "两条解析通道的文本输出汇合后，进入统一的切块处理阶段。这一步骤的目的是将长篇文本切割成若干适合检索的小片段，用于构建后续RAG流程所需的文本块库。"), "两路解析文本进入统一切块：PyMuPDF连续文本块按约1200字符组织，MinerU Markdown块为尽量保留表格行列关系按约1600字符组织，相邻块保留重叠区域。块大小描述的是索引单元，不等同于最终送入大语言模型的证据长度。", chapter3_body_style)
set_text(find_exact(doc, "切块的核心思路是将每页文本按固定长度滑动切割，相邻两块之间保留一段重叠区域。重叠设计是为了防止一句完整的表述恰好被切口截断，导致关键信息碎片化，进而影响后续检索的召回质量。PyMuPDF通道输出的是连续纯文本，块大小设为1200字符；MinerU通道输出的Markdown中包含完整的表格行，为避免表格被切割截断，块大小适当放宽至1600字符，重叠区域也相应加长。每个切块完成后被封装为一个标准化的文本块对象，记录其来源通道、所属页码以及文本内容，并赋予唯一编号以便后续溯源。全部文本块汇总构成文本块库，其结构可概括如表 5："), "每个文本块封装来源通道、原始页码、块编号和正文（表 4）。检索命中的证据在进入抽取提示前再按配置截取，当前单条上下文上限约1800字符。因此，1200、1600和1800分别对应两类索引块与抽取上下文上限，三者用途不同。", chapter3_body_style)
set_text(find_exact(doc, "这一动态混合解析策略，模型在保证核心信息高保真提取的同时，降低工程开销，提升模型在生产环境下的可扩展性。"), "查询改写依据字段标准名称、别名、必备证据词、领域词和期望单位生成不超过8条扩展查询。确定性改写不额外调用大语言模型，并将同一字段的检索线索约束在其Schema内。", chapter3_body_style)
replace_caption_title(find_exact(doc, "表 5  文本块库结构"), "文本块库结构")

set_text(find_exact(doc, "规则检索与语义检索"), "双路召回与本地回退", heading3_style)
set_text(find_exact(doc, "（1）规则检索"), "（1）查询改写与BM25召回", chapter3_body_style)
set_text(find_exact(doc, "检索层前置查询改写模块。对于每个待抽取字段，模型舍弃传统的单字段名检索模式，基于确定性规则进行多源线索扩展。改写来源涵盖字段别名集、必备证据词、预设领域词表，每个目标字段最终派生出不超过8条语义异构的扩展查询。该过程规避了传统方法对大语言模型的依赖，以极低的时间开销扩大了检索词汇的覆盖面，缓解了ESG异构表述导致的召回遗漏。"), "BM25路径接收改写后的多条查询，对专业术语、单位与表头短语进行稀疏检索。该路径适合定位“范围一”“tCO2e”等低频但区分度较高的线索，当前候选集合按配置截取。", chapter3_body_style)
set_text(find_exact(doc, "（2）改进BM25算法检索"), "（2）BM25得分", chapter3_body_style)
set_text(find_exact(doc, "在稀疏检索路径，模型基于改进BM25算法执行字段证据召回，无需维护外部向量索引，部署轻量。BM25算法的核心优势在于其对词频与文档稀缺性的精细建模。对于ESG报告这类专业文本，“碳排放”“范围一”“tCO2e”等指标词汇在全文中出现频次极低但信息价值极高，BM25的逆文档频率机制会自动为此类稀缺词赋予更高的检索权重，使包含这些术语的文本块在相关字段的检索中优先展示。另外，BM25算法内置的文档长度归一化惩罚机制，能够有效抑制长文本块通过堆砌普通词汇而虚高得分的情况，确保检索结果的质量不受文本长度干扰。BM25算法在本方案中的设置如下："), "BM25对词频、逆文档频率和文本块长度进行归一化，作为稀疏检索的基础得分，并叠加字段短语和期望单位等奖励。其计算形式如下：", chapter3_body_style)
set_text(find_exact(doc, "其中，第一项为基于BM25算法的词级统计路径，利用改进逆文档频率放大稀缺指标的权重，并基于全局平均长度进行归一化惩罚，抑制长文本块恶意堆砌词频；第二项为精准短语奖励路径，当文本块连续包含行业核心短语或特征单位时，直接触发硬编码奖励，确保高价值的数据表头或关键绩效区间能以绝对优势入选，最终输出得分最高的Top-K个文本块作为证据上下文。"), "式（6）中，q为字段查询，d为文本块，f(t,d)表示词项t在d中的频次，|d|与avgdl分别表示文本块长度和平均长度，k1、b为调节参数。字段短语、单位与必备词用于后续重排，不改变BM25基础公式。", chapter3_body_style)
set_text(find_exact(doc, "（3）自适应语义嵌入"), "（3）Embedding语义召回与本地回退", chapter3_body_style)
set_text(find_exact(doc, "自适应语义嵌入检索通道通过构建高效的文本向量索引，实现跨越字面表达的深层语义匹配。本方案设计了兼容异构环境的语义编码自适应切换系统。在线模式下，调度阿里云DashScope的text-embedding-v4高维模型，精准捕捉金融及环境治理语境下的细粒度表征；离线或轻量化部署时则退回至基于Blake2b哈希的384维n-gram局部嵌入，在零外部依赖下保障检索连续性。候选向量经归一化后统一为余弦相似度度量，并辅以签名缓存规避冗余编码开销。"), "Embedding路径使用配置的语义向量服务召回跨表述证据。接口不可用、请求失败或未配置时，系统回退到本地字符n-gram TF-IDF余弦检索，并在任务日志中记录回退状态。该回退能力维持检索流程连续性，但不等同于线上Embedding服务的语义效果，企业可据日志决定是否重跑。", chapter3_body_style)

set_text(find_exact(doc, "证据融合排序"), "RRF融合、字段重排与Top-K输出", heading3_style)
set_text(find_exact(doc, "为消除双路分数标准不同对后续检索的影响，本方案采用排序融合（RRF）算法对BM25算法结果与向量检索结果进行综合排序，公式如下："), "BM25、Embedding或本地回退路径的分数尺度不同，因此系统依据候选名次采用互惠排序融合（RRF）统一排序：", chapter3_body_style)
set_text(find_exact(doc, "其中为为检索路径集合，为文本块在不同路径中的名次。RRF有效提升了在多路检索中表现稳定的候选块排名。融合排序后，模型依据字段名、别名、必备词与禁用词特征对候选块进行打分，重新排序。命中设定特征的文本块获得额外权重溢价，而包含“指标索引”等宏观导航词的文本块则予以惩罚下调。通过二次重排确保高价值数据表头与关键绩效证据以绝对优势进入Top-K，最终输出至大语言模型进行结构化字段抽取。"), "式（7）中，R为检索路径集合，rank_r(d)为文本块d在路径r中的名次，k为平滑常数。融合后，系统依据目标字段名称、别名、必备词、期望单位和排除词进行字段级重排，并输出Top-K证据。字段专属奖励与排除约束可降低同一表格中相邻字段共用上下文造成的串扰。", chapter3_body_style)

# 3.4 extraction and quality.
set_text(find_exact(doc, "大模型字段提取"), "批量字段抽取", heading3_style)
set_text(find_exact(doc, "字段抽取环节采用批次化调度策略，将60个ESG指标按预设步长组织为动态任务批次，通过单一JSON对象响应完成多字段并发抽取，如图 6。"), "字段抽取采用批次化调度，不将“多字段同批处理”描述为并发。当前60个字段默认每批6个字段，每份报告最多12次大语言模型调用，每个字段使用检索阶段返回的Top-K证据，处理流程见图 6。", chapter3_body_style)
set_text(find_exact(doc, "模型先从RAG文本块库中为每个字段召回Top-K重排序证据，再将多组字段与其对应上下文合并为结构化提示输入大语言模型。提示词设计遵循高鲁棒性原则，设定的系统提示词明确约束模型仅基于给定证据抽取，禁止编造信息。每个字段经压缩后保留核心元数据，证据片段截取前1200字符并标注来源页码与得分，输出格式强制约束为JSON，温度参数设为0消除随机性。对于未披露或证据不足的字段，模型通过标志位实现熔断，避免强行输出无效值。定量字段优先映射值、单位、年份，定性字段输出语义摘要，实现定量与定性分流处理。"), "系统将同批字段及其证据组织为结构化提示，约束大语言模型只依据召回证据输出JSON。定量字段返回值、单位和年份，定性字段返回证据摘要；证据不足或报告未披露时返回未命中状态。每条结果同时记录来源页码和证据，供后处理与人工复核使用。", chapter3_body_style)
set_text(find_exact(doc, "目标年份对齐 "), "目标年份对齐", heading3_style)
set_text(find_exact(doc, "RAG检索虽然能够召回语义相关证据，但若缺乏时间约束，容易将ESG报告同时披露的历史年份数据误作为最终结果。因此，本方案在后处理阶段引入目标年份对齐机制，对抽取结果进行时间一致性校验，提高结构化数据的准确性。"), "ESG报告常同时列示多个年度数据，系统在后处理阶段校验字段值与目标年份的对应关系。当前配置默认以2024年为目标年份，不能据此宣称对任意年度无需配置即可适配。", chapter3_body_style)
set_text(find_exact(doc, "模型首先根据报告文件名、首页标题及发布日期确定目标统计年份，并作为后续校验的统一依据。随后，对候选证据执行年份识别，通过正则表达式提取时间信息，并结合上下文建立年份与字段值之间的对应关系，避免将发布日期或历史事件年份误识别为指标统计年份。"), "系统从证据文本中识别年份信息，并将字段结果与配置的目标年份比较。报告文件名、首页标题和发布日期可作为人工复核线索，但最终校验以当前任务配置为准，避免把披露日期或历史事件年份误作统计年份。", chapter3_body_style)
set_text(find_exact(doc, "在一致性校验阶段，模型比较目标年份与证据年份。当二者一致时，直接保留抽取结果；若存在年份偏移，则优先从候选证据中重新选择目标年份对应的数据。当未检索到符合年份约束的证据时，模型保留当前结果并降低其置信度，同时标记为“待复核”，进入人工审核流程。"), "年份一致时保留结果；存在年份偏移、年份缺失或跨年度表格定位不清时，系统标记年份异常并降低复核优先级中的可信等级，不把不确定结果直接作为目标年度数据交付。", chapter3_body_style)
set_text(find_exact(doc, "此外，针对ESG报告中大量存在的跨年度对比表格，模型结合表格结构定位目标年份所在列，并提取对应数值，减少相邻年份数据误匹配的问题。通过目标年份对齐机制，模型能够有效保证字段值与目标统计年度保持一致，提高ESG指标结构化结果的时效性、准确性与可比性。"), "对跨年度表格，系统结合证据中的年份与数值邻近关系进行校验；无法建立清晰对应关系时进入人工复核。该机制用于暴露年份风险，而不承诺自动消除所有跨年误配。", chapter3_body_style)

quality_heading = find_exact(doc, "3.4.3 抽取结果质量校验")
set_text(quality_heading, "证据链评分与复核优先级", heading3_style)
set_text(find_exact(doc, "本方案还加入了数据质量检测指标，引入置信度评估机制，采用归一化纠偏策略，当模型输出置信度溢出预设边界时，自动缩放至0-1区间，保证不同字段间置信度具有统一的量化标准和可比性。证据链得分模型从字段名匹配、数值命中、单位命中、年份命中、来源页码存在性及证据完整性六个维度进行加权评分，总分100分。得分低于50分的记录自动标记为弱证据链，并生成相应提示信息，引导人工优先复核低可信度结果，从而进一步提升最终数据质量。"), "系统将模型置信度归一化至0—1区间，并从字段名匹配、数值命中、单位命中、年份命中、页码存在性和证据完整性六个维度计算证据链得分。置信度低于0.7或证据链得分低于50的字段进入重点复核；单位异常、年份异常和缺少页码分别保留质量标记，便于用户按风险排序处理。", chapter3_body_style)

budget_para = find_exact(doc, "为控制成本与保障模型稳定性，模型设置报告级LLM调用次数硬上限。当累计调用频次逼近阈值时，自动终止后续批次请求，对剩余字段执行空结果降级填充，并标注预算耗尽原因，避免因报告异常导致计算资源过度消耗。同时，模型保留完整的任务执行日志和异常状态信息，便于后续定位问题与重新处理，将单份报告处理成本约束在可预期边界内，从而保障效率、资源消耗与模型稳定性之间的平衡，为大规模ESG报告批处理场景提供了可量化的服务质量保障边界。")
insert_before(budget_para, "调用预算与异常降级", heading3_style)
set_text(budget_para, "系统将单份报告的大语言模型调用上限设置为可配置的资源使用边界，当前默认最多12次。预算耗尽时，剩余字段返回带原因的空结果；调用失败时记录错误并允许任务重跑。解析工具失败、Embedding接口失败、缺少页码、单位或年份异常均保留对应状态，避免异常被静默写入正常结果。", chapter3_body_style)

# 3.5 metrics and enterprise readiness.
set_text(find_exact(doc, "模型性能指标"), "模型性能与工程可用性验证")
set_text(find_exact(doc, "为验证模型在实际ESG报告处理任务中的提取能力，本方案采用人工标注对输出结果进行逐字段核验。本次评估共抽检核验共100项相关指标。其中16项经人工确认为报告未披露，不计入精度评估，最终有效样本为84项，漏抽与待定结果均为0，说明模型对有效样本实现了全覆盖输出，无遗漏。"), "本次评估采用人工逐字段核验，共检查100项，其中16项经人工确认原报告未披露，84项为已披露有效样本。84项按互斥类别统计：56项自动正确且可直接采纳，12项结果正确但被质量机制标记为待复核，12项字段串扰，3项数值错误，1项年份错误，数量合计为84项。", chapter3_body_style)
replace_caption_title(find_exact(doc, "图 7  模型性能指标图"), "结构化提取人工核验结果")
set_text(find_exact(doc, "评估结果显示（图 7），84项有效样本中模型直接输出正确且被人工采纳的结果为56项，采纳正确率为66.67%；另外有12项输出结果经人工核查确认正确，但因置信度偏低或证据链不足被平台预警为待复核，实际可用于分析的正确结果共68项，人工确认后准确率达80.95%。两项指标之间差异说明，模型内置的质量校验机制确实将一部分可用结果标记为了保守状态，这在复核资源充足的场景下有助于减少误用风险，但也意味着约14.29%的正确结果需要人工确认后才能释放。"), "如图 7所示，56项自动正确结果占84项有效样本的66.67%，统一称为“自动直接采纳率”；再加上12项经人工确认正确的待复核结果，可用结果共68项，占80.95%，统一称为“人工确认后可用率”。80.95%不表述为模型准确率，其含义是质量标记与人工确认共同作用后的可用结果比例。", chapter3_body_style)
set_text(find_exact(doc, "从错误分布来看，字段串扰是当前模型最主要的误差来源，共12项，占有效样本的14.29%。字段串扰的典型表现为模型召回了目标字段相邻指标的证据，并以较高置信度完成了抽取，导致值和单位均合规，但字段归属错误。这类错误在ESG报告中具有一定普遍性，根本原因在于同一数据表内相邻行列的指标在文本层面高度相似，BM25检索与向量检索均难以在词汇层面实现精确区分。此外，值错误3项，占比3.57%，主要集中于数值格式不统一或单位换算判断有误的情形；年份错误1项，占比1.19%，因跨年度对比表中目标年份定位偏移所致。"), "错误中，字段串扰12项，占有效样本14.29%，主要表现为同表相邻行列的证据被映射到错误字段；数值错误3项，占3.57%；年份错误1项，占1.19%。本次小样本未观测到漏抽，但该观察不能外推为系统无漏抽，也不能替代更大规模、跨行业和跨版式测试。", chapter3_body_style)
set_text(find_exact(doc, "综合来看，模型已能够满足金融领域严肃文本处理中“自动提取为主、人工复核为辅”的工程定位。ESG报告作为重要的金融监管与投资决策依据，对数据准确性和可追溯性要求极高，本模型通过RAG证据召回、质量校验与复核机制，确保每项指标均有明确来源支撑。在实际工程应用中，多数字段可直接输出可用结果，少量需人工干预的字段特征明确，便于复核人员快速定位与确认，显著提升整体处理效率。模型未出现系统性漏抽，验证了召回机制的性能完整可靠。"), "该结果表明，当前系统可将一部分字段直接交付，并通过质量标记将另一部分结果交由人工确认；字段串扰仍是主要改进方向。评估样本规模较小，且16项“未披露”只由人工确认后从值准确率中剔除，尚未对系统识别未披露状态的准确性进行独立评价。", chapter3_body_style)
set_text(find_exact(doc, "在工程落地层面，本方案设计为“开箱即用”形态。用户只需登录本方案开发的ESGMiner平台，上传需要提取的原始PDF，即可端到端地获得处理后的数据。ESGMiner平台支持批量PDF上传、异步任务处理、字段级复核界面与一键导出CSV，结合清晰的证据链和置信度标注，极大降低了人工干预成本，可直接服务于监管审查、投资分析及企业ESG数据治理场景，提升自动化程度，为用户显著降低人工复核成本、加快处理流程并提供更稳定可靠的ESG数据输出。"), "工程可用性不仅取决于正常路径，还取决于输入异常、外部服务失败和资源预算耗尽时是否给出可识别结果。表 5汇总现有企业运行能力及降级策略。本文所称“配置式部署、快速上线”是指：企业完成运行环境和模型服务参数配置后，可通过标准化部署方式启动系统，并直接使用报告上传、任务处理、字段复核、结果导出和指标分析功能。", chapter3_body_style)

ch4 = find_exact(doc, "ESGMiner平台设计与实现")
table_caption = add_seq_caption_before(ch4, "表", 5, "企业运行能力与降级策略", caption_style)
run_table = doc.add_table(rows=10, cols=4)
run_table.alignment = WD_TABLE_ALIGNMENT.CENTER
run_table.autofit = False
if doc.tables[3].style is not None:
    run_table.style = doc.tables[3].style
widths = [Cm(3.0), Cm(3.4), Cm(5.2), Cm(4.2)]
headers = ["运行场景", "系统能力", "异常或边界", "用户获得的结果"]
data = [
    ("非目标报告", "规则准入与主动跳过", "未命中正向词或命中排除词", "跳过状态与原因代码"),
    ("重复文件", "SHA-256识别", "历史任务存在相同文件指纹", "冲突提示与既有任务信息"),
    ("批量处理", "批量上传与异步任务", "单个任务失败不替代其他任务结果", "独立任务编号与状态查询"),
    ("复杂页面", "MinerU选择性调用", "页数上限、未启用或调用失败", "基础文本、路由及降级状态"),
    ("语义检索", "Embedding接口", "接口失败或未配置", "本地字符n-gram TF-IDF回退记录"),
    ("字段抽取", "报告级LLM调用预算", "预算耗尽或模型调用失败", "带原因的空结果及重跑依据"),
    ("质量控制", "低可信字段重点复核", "置信度＜0.7或证据链＜50", "待复核标记、证据与页码"),
    ("结果交付", "API返回JSON，导出中心提供CSV", "仅完成任务可导出", "系统集成或离线分析数据"),
    ("部署上线", "Docker Compose配置式部署", "需配置运行环境与模型服务参数", "标准化启动与平台功能入口"),
]
for j, h in enumerate(headers):
    cell = run_table.cell(0, j)
    set_cell_text(cell, h, bold=True, color=(255, 255, 255), size=8.5, align=WD_ALIGN_PARAGRAPH.CENTER)
    shade_cell(cell, "2F7D64")
    cell.width = widths[j]
for i, row in enumerate(data, 1):
    for j, value in enumerate(row):
        cell = run_table.cell(i, j)
        set_cell_text(cell, value, size=8.2, align=WD_ALIGN_PARAGRAPH.CENTER if j < 2 else WD_ALIGN_PARAGRAPH.LEFT)
        shade_cell(cell, "EEF6F2" if i % 2 else "FFFFFF")
        cell.width = widths[j]

# Move the newly appended table to immediately after its caption.
run_table._tbl.getparent().remove(run_table._tbl)
table_caption._p.addnext(run_table._tbl)
transition = insert_before(ch4, "上述算法与工程机制完成了从PDF输入到结构化结果、质量标记和复核数据的后端闭环。下一章将进一步说明这些能力如何通过ESGMiner平台转化为面向非技术用户的报告管理、结果复核、指标分析与数据导出功能。", chapter3_body_style)

# Update references whose visible text is not dynamically regenerated until Word refreshes fields.
for p in doc.paragraphs:
    if p.text.strip() == "表 3  评分规则与阈值":
        replace_caption_title(p, "评分规则与阈值")
    elif p.text.strip() == "表 4  各区间处理策略":
        replace_caption_title(p, "各区间处理策略")

# Keep original and revised files separate; request Word to refresh all fields on open.
settings = doc.settings._element
update_fields = settings.find(qn("w:updateFields"))
if update_fields is None:
    update_fields = OxmlElement("w:updateFields")
    settings.append(update_fields)
update_fields.set(qn("w:val"), "true")

doc.save(OUT)
print(OUT)
print(FLOW_PNG)
