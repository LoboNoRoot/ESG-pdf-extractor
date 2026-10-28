from pathlib import Path
from docx import Document

root = Path(r"C:\Users\Lenovo\Desktop\省一改")
path = root / "项目报告书-智析万卷：面向ESG报告的自适应布局分析与要素提取-专业团队ProMax-国赛专项修订版.docx"
doc = Document(path)


def set_seq_result(paragraph, number: int):
    after_sep = False
    for run in paragraph.runs:
        fld = run._r.xpath(".//w:fldChar/@w:fldCharType")
        if "separate" in fld:
            after_sep = True
            continue
        if "end" in fld and after_sep:
            break
        if after_sep and run.text:
            run.text = str(number)
            after_sep = False


def replace_visible_run(paragraph, old: str, new: str):
    for run in paragraph.runs:
        if old in run.text:
            run.text = run.text.replace(old, new)
            return
    raise RuntimeError(f"Visible run not found: {old!r} in {paragraph.text!r}")


# Update cached SEQ results without removing Word fields.
table_captions = []
for p in doc.paragraphs:
    instr = "".join(p._p.xpath(".//w:instrText/text()"))
    if "SEQ 表" in instr:
        table_captions.append(p)
for number, p in enumerate(table_captions, 1):
    set_seq_result(p, number)

# Update cached contents of the main TOC and figure/table directory.
toc_title = next(p for p in doc.paragraphs if p.text.startswith("3.5 模型性能指标"))
replace_visible_run(toc_title, "模型性能指标", "模型性能与工程可用性验证")

fig4 = next(p for p in doc.paragraphs[:70] if p.text.startswith("图 4  分流决策核心分支图"))
replace_visible_run(fig4, "分流决策核心分支图", "企业端到端交付闭环图")
fig7 = next(p for p in doc.paragraphs[:70] if p.text.startswith("图 7  模型性能指标图"))
replace_visible_run(fig7, "模型性能指标图", "结构化提取人工核验结果")

table_entries = [p for p in doc.paragraphs[:70] if p.text.startswith("表 ") and "\t" in p.text]
desired = [
    (" 1  ", "性能评估口径说明"),
    (" 2  ", "评分规则与阈值"),
    (" 3  ", "各区间处理策略"),
    (" 4  ", "文本块库结构"),
    (" 5  ", "企业运行能力与降级策略"),
]
if len(table_entries) != 5:
    raise RuntimeError(f"Expected 5 table-directory entries, got {len(table_entries)}")
for p, (num_text, title) in zip(table_entries, desired):
    p.runs[1].text = num_text
    p.runs[2].text = title

doc.save(path)
print(path)
print("TABLE_CAPTIONS", [p.text for p in table_captions])
print("TABLE_DIRECTORY", [p.text for p in table_entries])
